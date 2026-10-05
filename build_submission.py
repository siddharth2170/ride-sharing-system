from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)


def font(size=28, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Menlo Bold.ttf" if bold else "/System/Library/Fonts/Menlo.ttc",
        "/System/Library/Fonts/Supplemental/Courier New Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Courier New.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def code_image(source, dest, start, end, title):
    lines = Path(source).read_text().splitlines()[start - 1:end]
    width, line_h = 1900, 33
    height = 90 + line_h * len(lines) + 35
    im = Image.new("RGB", (width, height), "#F7F7F8")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, width, 64), fill="#243447")
    d.text((28, 16), title, font=font(26, True), fill="white")
    body = font(23)
    number = font(21)
    colors = {"class": "#7A3E9D", "virtual": "#006D77", "override": "#006D77",
              "private:": "#B45309", "public:": "#B45309", "protected:": "#B45309",
              "subclass:": "#7A3E9D", "methodsFor:": "#006D77"}
    y = 78
    for i, line in enumerate(lines, start):
        d.text((24, y), f"{i:>3}", font=number, fill="#89929B")
        color = "#1F2933"
        for token, candidate in colors.items():
            if token in line:
                color = candidate
                break
        d.text((95, y), line.expandtabs(4), font=body, fill=color)
        y += line_h
    im.save(dest)


def terminal_image(text_path, dest, title):
    lines = Path(text_path).read_text().splitlines()
    width, line_h = 1900, 38
    height = 105 + line_h * len(lines) + 35
    im = Image.new("RGB", (width, height), "#111827")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, width, 64), fill="#263244")
    for x, c in [(28, "#FF5F56"), (62, "#FFBD2E"), (96, "#27C93F")]:
        d.ellipse((x, 21, x + 20, 41), fill=c)
    d.text((145, 15), title, font=font(25, True), fill="#E5E7EB")
    y = 82
    body = font(25)
    for line in lines:
        color = "#7DD3FC" if line.startswith("===") or line.startswith("Polymorphic") else "#E5E7EB"
        d.text((32, y), line, font=body, fill=color)
        y += line_h
    im.save(dest)


code_image(ROOT / "cpp/main.cpp", ASSETS / "cpp-code.png", 1, 51,
           "C++ source excerpt — base class and overridden fare methods")
code_image(ROOT / "smalltalk/RideSharing.st", ASSETS / "smalltalk-code.png", 1, 63,
           "Smalltalk source excerpt — inheritance and message polymorphism")
terminal_image(ASSETS / "cpp-output.txt", ASSETS / "cpp-output.png", "C++ compiled program output")
terminal_image(ASSETS / "smalltalk-output.txt", ASSETS / "smalltalk-output.png", "Smalltalk Transcript output")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"; normal.font.size = Pt(12); normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 2
normal.paragraph_format.space_after = Pt(0)
for name in ["Title", "Heading 1", "Heading 2"]:
    st = styles[name]
    st.font.name = "Times New Roman"; st.font.color.rgb = RGBColor(0, 0, 0); st.font.bold = True
styles["Title"].font.size = Pt(14)
styles["Heading 1"].font.size = Pt(12)
styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
styles["Heading 1"].paragraph_format.space_before = Pt(0)
styles["Heading 1"].paragraph_format.space_after = Pt(0)
styles["Heading 2"].font.size = Pt(12); styles["Heading 2"].font.italic = True
styles["Heading 2"].paragraph_format.space_before = Pt(0); styles["Heading 2"].paragraph_format.space_after = Pt(0)

# Remove any theme-provided Title border so the APA title remains plain black text.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
caption.font.name = "Times New Roman"; caption.font.size = Pt(11); caption.font.color.rgb = RGBColor(0,0,0)
caption.paragraph_format.line_spacing = 1
caption.paragraph_format.space_after = Pt(6)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run()
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)

for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Object Oriented Ride Sharing System in Smalltalk and C++")
for text in ["Student Name", "Institution", "Course Number and Name", "Instructor Name", "October 4, 2026"]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text)

doc.add_page_break()
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Object Oriented Ride Sharing System in Smalltalk and C++")

def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Inches(.5)
    return p

body("This project implements the same ride sharing domain in C++ and Smalltalk to demonstrate encapsulation, inheritance, and polymorphism. Both programs model standard and premium rides, a driver’s completed rides, and a rider’s request history. Mixed ride objects are processed through one interface, while each subtype calculates its own fare.")

doc.add_heading("System Design", level=1)
body("Ride stores an identifier, pickup and drop-off locations, and distance. StandardRide charges a $3.00 booking fee plus $2.00 per mile; PremiumRide charges $7.50 plus $3.50 per mile. Driver and Rider record completed and requested rides through methods. Validation rejects negative distance and invalid ratings.")

doc.add_heading("Encapsulation", level=1)
body("In C++, fields are private, including Driver::assignedRides_ and Rider::requestedRides_. Clients add rides only through public methods, leaving each class in control of its collection. In Smalltalk, instance variables are scoped to the class and its methods; clients send addRide: or viewRides rather than manipulating collections. Both versions therefore hide representation details behind a focused public protocol.")

doc.add_heading("Inheritance", level=1)
body("StandardRide and PremiumRide inherit identity, route, distance, and display behavior from Ride. Each supplies only rideType and fare. C++ declares these as pure virtual functions and shares a protected helper. Smalltalk uses subclassResponsibility and shares baseFareAt:bookingFee:. The hierarchy removes duplicated state and supports additional ride categories (Black et al., 2009; Stroustrup, 2013).")

doc.add_heading("Polymorphism", level=1)
body("Each demonstration places standard and premium objects in one collection. C++ stores shared pointers to Ride and dispatches virtual calls at runtime. Smalltalk sends identical messages to heterogeneous objects, and dynamic lookup selects the subtype method. A six-mile premium trip produces $28.50; a 12.5-mile standard trip produces $28.00. No type test is needed.")

doc.add_heading("Language Comparison", level=1)
body("C++ makes the boundary explicit with access specifiers, const methods, pure virtual functions, override, and smart pointers; the compiler checks much of the interface. Smalltalk uses fewer declarations because objects respond dynamically to messages. Despite different type systems, both versions bind data to behavior, specialize Ride through inheritance, and permit substitutable ride objects.")

doc.add_heading("Verification and Repository", level=1)
body("The C++ program compiled with C++17, -Wall, -Wextra, and -pedantic and executed successfully (Appendix B). The Smalltalk source uses Pharo/Squeak file-in syntax and includes a Transcript demonstration (Appendix D). The complete source and documentation are available in the public repository below.")
p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.5)
r = p.add_run("GitHub repository: "); r.bold = True
p.add_run("https://github.com/siddharth2170/ride-sharing-system")

doc.add_heading("Conclusion", level=1)
body("Both implementations meet the required functionality. Encapsulation protects ride histories, inheritance centralizes shared behavior, and polymorphism selects subtype-specific fares through a common Ride interface. The result is compact and extensible in both languages.")

doc.add_page_break()
doc.add_heading("References", level=1)
refs = [
    "Black, A. P., Ducasse, S., Nierstrasz, O., Pollet, D., Cassou, D., & Denker, M. (2009). Pharo by example. Square Bracket Associates.",
    "Stroustrup, B. (2013). The C++ programming language (4th ed.). Addison-Wesley Professional.",
]
for t in refs:
    p = doc.add_paragraph(t)
    p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.5)

def appendix(title, fig_no, caption_text, image_path):
    doc.add_page_break()
    doc.add_heading(title, level=1)
    p = doc.add_paragraph(style="Figure Caption"); p.add_run(f"Figure {fig_no}").bold = True
    p = doc.add_paragraph(style="Figure Caption"); p.add_run(caption_text).italic = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(6.45))

appendix("Appendix A C++ Code", 1, "Ride hierarchy and overridden fare methods in C++", ASSETS / "cpp-code.png")
appendix("Appendix B C++ Sample Output", 2, "Compiled C++ program output", ASSETS / "cpp-output.png")
appendix("Appendix C Smalltalk Code", 3, "Ride hierarchy and overridden fare methods in Smalltalk", ASSETS / "smalltalk-code.png")
appendix("Appendix D Smalltalk Sample Output", 4, "Smalltalk Transcript output for the demonstration", ASSETS / "smalltalk-output.png")

props = doc.core_properties
props.title = "Object Oriented Ride Sharing System in Smalltalk and C++"
props.subject = "Encapsulation, inheritance, and polymorphism"
props.author = "Student Name"

doc.save(OUT / "Ride_Sharing_System_APA7_Report.docx")
print(OUT / "Ride_Sharing_System_APA7_Report.docx")
